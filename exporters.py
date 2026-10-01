"""Export customer tickets to CSV, Excel (xlsx) and PDF."""

import csv
import io
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def build_rows(customer, tickets, custom_fields):
    header = ["Ticket ID", "Customer Name", "Status", "Priority",
              "Created Date", "Last Updated", "Updated Notes"]
    header += [f.name for f in custom_fields]
    rows = []
    for t in tickets:
        data = json.loads(t.custom_data or "{}")
        row = [
            t.id, customer.name, t.status or "Open", t.priority or "Medium",
            t.created_date.strftime("%Y-%m-%d %H:%M") if t.created_date else "",
            t.updated_date.strftime("%Y-%m-%d %H:%M") if t.updated_date else "",
            t.updated_notes or "",
        ]
        row += [data.get(f.name, "") for f in custom_fields]
        rows.append(row)
    return header, rows


def to_csv(header, rows):
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(header)
    writer.writerows(rows)
    # BOM so Excel opens UTF-8 CSV correctly
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def to_excel(header, rows, customer_name):
    wb = Workbook()
    ws = wb.active
    ws.title = "Tickets"[:31]
    ws.append(header)
    for r in rows:
        ws.append(r)

    head_fill = PatternFill("solid", fgColor="4F46E5")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for i, col in enumerate(header, start=1):
        longest = max([len(str(col))] + [len(str(r[i - 1])) for r in rows]) if rows else len(col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(12, longest + 2), 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(header, rows, customer_name):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm,
                            title=f"{customer_name} - Tickets")
    styles = getSampleStyleSheet()
    cell_style = styles["BodyText"]
    cell_style.fontSize = 8
    cell_style.leading = 10

    story = [
        Paragraph(f"Ticket Report &mdash; {customer_name}", styles["Title"]),
        Paragraph(f"Total tickets: {len(rows)}", styles["Normal"]),
        Spacer(1, 6 * mm),
    ]

    def esc(v):
        return str(v).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    data = [[Paragraph(f'<font color="white"><b>{esc(h)}</b></font>', cell_style) for h in header]]
    data += [[Paragraph(esc(v), cell_style) for v in r] for r in rows]
    if not rows:
        data.append([Paragraph("No tickets", cell_style)] + [""] * (len(header) - 1))

    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D1D5DB")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(table)
    doc.build(story)
    return buf.getvalue()
